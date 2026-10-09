"""Build the SIH 2026 idea deck on the official SIH template (6 slides, NexaBuild).

    python -I tools/build_sih_deck.py            # writes docs/NexaBuild_VAYU-SARTHI_SIH2026.pptx
    python -I tools/build_sih_deck.py --live https://<user>-vayu-sarthi.hf.space   # adds the live link

The template's frame is kept as it is: the SIH title and logo, the team badge, the slide titles, the
footer bar and the idea-detail pointers (used word for word as section labels). The instructions slide is
removed, as the template allows. Screenshots come from the running prototype
(tools/deck_assets/capture.mjs); numbers come from the benchmarks in docs/PLAN.md §8.
Convert to PDF for the portal:  soffice --headless --convert-to pdf docs/NexaBuild_VAYU-SARTHI_SIH2026.pptx
"""
from __future__ import annotations

import argparse
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "tools" / "deck_assets"
OUT = ROOT / "docs" / "NexaBuild_VAYU-SARTHI_SIH2026.pptx"
REPO_URL = "https://github.com/DeepakSinghhh/Air-power-Dynamic"
LIVE_URL = ""  # set with --live (docs/DEPLOY.md)

TEAM = "NexaBuild"
TEAM_ID = ""  # left blank until the portal issues it

# Palette: the SIH logo's saffron and green on the template's navy and white. Saffron marks steps and
# callouts; green marks measured gains; everything else is navy and grey.
INK = RGBColor(0x1C, 0x2E, 0x4A)
TX2 = RGBColor(0x1F, 0x49, 0x7D)    # the template's own heading navy
SAFF = RGBColor(0xC2, 0x56, 0x1A)
SAFF_TINT = RGBColor(0xFB, 0xEA, 0xDE)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
MUTED = RGBColor(0x55, 0x5F, 0x6E)
PANEL = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD3, 0xD9, 0xE1)
BASE_GREY = RGBColor(0xA9, 0xB1, 0xBD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


# ---------- helpers ----------

def _run(p, text, size, color=INK, bold=False, italic=False, font=FONT, link=None, underline=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    if underline:
        f.underline = True
    if link:
        r.hyperlink.address = link
        f.color.rgb = color
    return r


def _bullet(p, color=SAFF, indent=0.17, font="Arial", char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(color))
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", font)
    if font == "Wingdings":
        bf.set("pitchFamily", "2")
        bf.set("charset", "2")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def _pointer_bullet(p, color=SAFF, indent=0.3):
    """The template's own ❖ marker: a Wingdings "v" bullet."""
    _bullet(p, color, indent, font="Wingdings", char="v")


def text(slide, x, y, w, h, paras, size=12, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name=None, margin=0.0, after=0, line=None):
    """paras: list of paragraphs; a paragraph is a str, or a dict(runs=[(text, opts)], bullet=bool, after=pt)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        spec = {"runs": [(para, {})]} if isinstance(para, str) else para
        p.alignment = spec.get("align", align)
        p.space_after = Pt(spec.get("after", after))
        if line or spec.get("line"):
            p.line_spacing = spec.get("line", line)
        if spec.get("bullet"):
            _bullet(p)
        for t, o in spec["runs"]:
            _run(p, t, o.get("size", spec.get("size", size)), o.get("color", spec.get("color", color)),
                 o.get("bold", False), o.get("italic", False), o.get("font", FONT), o.get("link"))
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius / min(w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def badge(slide, x, y, label, d=0.32, fill=SAFF, color=WHITE, size=12, outline=None):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    if outline is not None:
        c.line.color.rgb = outline
        c.line.width = Pt(1.5)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, label, size, color, bold=True)
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5, head=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def pointer(slide, x, y, w, label, size=14, underline=False, h=0.32):
    """An idea-detail pointer from the template, word for word, in the template's own style."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    p = tf.paragraphs[0]
    _pointer_bullet(p, indent=0.3 if size < 16 else 0.36)
    r = _run(p, label, size, TX2, bold=True)
    if underline:
        r.font.underline = True
    return tb


def picture(slide, path, x, y, w=None, h=None, border=RULE):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(0.75)
    return pic


def table(slide, x, y, widths, rows, header_fill=INK, size=11, row_h=0.4, header_h=0.4, zebra=True,
          header_pointers=False):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                                   Inches(header_h + row_h * (len(rows) - 1)))
    tbl = shape.table
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(header_h if i == 0 else row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (PANEL if zebra and i % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            runs = val if isinstance(val, list) else [(val, {})]
            p = tf.paragraphs[0]
            if i == 0 and header_pointers:
                _pointer_bullet(p, color=RGBColor(0xF3, 0xB0, 0x83), indent=0.24)
            for t, o in runs:
                _run(p, t, o.get("size", size), WHITE if i == 0 else o.get("color", INK),
                     bold=(i == 0) or o.get("bold", False), italic=o.get("italic", False),
                     font=o.get("font", FONT), link=o.get("link"))
            tcPr = cell._tc.get_or_add_tcPr()
            ln = etree.Element(qn("a:lnB"), w="6350")
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", str(RULE))
            tcPr.insert(0, ln)
    return tbl


def frame(slide, title=None):
    """Template frame: team badge, idea title; drop the template's prompt text box."""
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
            p = sh.text_frame.paragraphs[0]
            for r in list(p.runs)[1:]:
                p._p.remove(r._r)
            p.runs[0].text = TEAM
            f = p.runs[0].font
            f.size, f.bold, f.name = Pt(13), True, FONT
            f.color.rgb = INK
            tf = sh.text_frame
            tf.word_wrap = False
            for side in ("left", "right"):
                setattr(tf, f"margin_{side}", 0)
            sh.line.color.rgb = TX2
        elif sh.name.startswith("TextBox 8"):
            sh._element.getparent().remove(sh._element)
        elif title and sh.is_placeholder and sh.placeholder_format.type is not None \
                and "TITLE" in str(sh.placeholder_format.type):
            for r in sh.text_frame.paragraphs[0].runs:
                if r.text.strip():
                    r.text = title


# ---------- slides ----------

def slide_title(s):
    for sh in list(s.shapes):
        if sh.name == "Subtitle 3":
            sh.left, sh.top, sh.width, sh.height = Inches(0.42), Inches(1.12), Inches(5.7), Inches(1.35)
            tf = sh.text_frame
            tf.word_wrap = True
            for p in list(tf.paragraphs)[1:]:
                tf._txBody.remove(p._p)
            p = tf.paragraphs[0]
            for r in list(p.runs):
                p._p.remove(r._r)
            p.alignment = PP_ALIGN.LEFT
            _run(p, "VAYU-SARTHI", 34, TX2, bold=True, font="Times New Roman")
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.space_before = Pt(2)
            _run(p2, "Plans the air tasking day, and re-plans it in seconds when something changes, "
                     "moving as few aircraft as possible.", 14, MUTED)
        elif sh.name == "TextBox 9":
            sh.left, sh.top, sh.width = Inches(0.36), Inches(2.62), Inches(5.75)
            fields = [
                ("Problem Statement ID – ", "26250"),
                ("Problem Statement Title – ", "Air Power: Dynamic Air Operations & Resource Optimisation"),
                ("Theme – ", "Transportation & Logistics"),
                ("PS Category – ", "Software"),
                ("Team ID – ", TEAM_ID),
                ("Team Name – ", TEAM),
            ]
            paras = [p for p in sh.text_frame.paragraphs if p.runs]
            for p, (label, value) in zip(paras, fields):
                p.line_spacing = 1.25
                p.space_after = Pt(9)
                p.alignment = PP_ALIGN.LEFT
                p.runs[0].text = label
                p.runs[0].font.size = Pt(17)
                p.runs[0].font.color.rgb = INK
                if value:
                    _run(p, value, 17, TX2 if label.startswith("Team Name") else INK,
                         bold=label.startswith("Team Name"))
            first = sh.text_frame.paragraphs[0]
            if not first.runs:  # the template's empty first line
                sh.text_frame._txBody.remove(first._p)
    if LIVE_URL:
        text(s, 0.42, 6.62, 5.7, 0.36, [{"runs": [
            ("Live prototype: ", {"bold": True, "color": GREEN}),
            (LIVE_URL.removeprefix("https://"), {"link": LIVE_URL, "color": TX2, "bold": True})]}], size=14)


def slide_solution(s):
    frame(s, "VAYU-SARTHI")
    pointer(s, 0.45, 1.2, 12.4, "Proposed Solution (Describe your Idea/Solution/Prototype)", size=18,
            underline=True, h=0.4)

    # Left: the three pointers.
    lx, lw = 0.45, 6.15
    badge(s, lx, 1.8, "1")
    text(s, lx + 0.45, 1.81, lw - 0.45, 0.3, [{"runs": [("Detailed explanation of the proposed solution",
                                                         {"bold": True})]}], size=13.5)
    text(s, lx + 0.45, 2.14, lw - 0.45, 1.1, [
        {"runs": [("One live picture: aircraft, crews, weapons, weather, threats", {})], "bullet": True, "after": 3},
        {"runs": [("Builds the full air tasking plan for the day in under 10 s", {})], "bullet": True, "after": 3},
        {"runs": [("When something changes, proposes the ", {}), ("smallest fix", {"bold": True}),
                  (" as a diff", {})], "bullet": True, "after": 3},
        {"runs": [("The commander approves or rejects; nothing moves on its own", {})], "bullet": True},
    ], size=12)

    badge(s, lx, 3.36, "2")
    text(s, lx + 0.45, 3.37, lw - 0.45, 0.3, [{"runs": [("How it addresses the problem", {"bold": True})]}],
         size=13.5)
    table(s, lx + 0.45, 3.72, [1.82, 3.88], [
        ["The problem asks for", "What VAYU-SARTHI does"],
        [[("Dynamic retasking", {"bold": True})], "Re-plans in ~3 s; the rest of the plan stays put"],
        [[("Resource optimisation", {"bold": True})], "One model: aircraft, crew rest, weapons, tankers, reserve"],
        [[("Weather, uncertainty", {"bold": True})], "Fog forecast from real airport data; 2,000 simulated days"],
        [[("Decision support", {"bold": True})], "Gives a reason for each choice; a human approves it"],
    ], size=10.5, row_h=0.3, header_h=0.3)

    badge(s, lx, 5.45, "3")
    text(s, lx + 0.45, 5.46, lw - 0.45, 0.3, [{"runs": [("Innovation and uniqueness of the solution",
                                                         {"bold": True})]}], size=13.5)
    text(s, lx + 0.45, 5.79, lw - 0.45, 1.0, [
        {"runs": [("Changes as little as possible: a diff, not a brand-new plan", {})], "bullet": True, "after": 2},
        {"runs": [("Says what it would take to fit any mission it had to drop", {})], "bullet": True, "after": 2},
        {"runs": [("Plain-language copilot that answers only from the engine", {})], "bullet": True, "after": 2},
        {"runs": [("Runs offline on one laptop; also plans flood and earthquake relief", {})], "bullet": True},
    ], size=12)

    # Right: the working prototype, annotated.
    ix, iy, iw = 6.85, 1.8, 6.03
    ih = iw * 1632 / 2620
    picture(s, ASSETS / "retask.jpg", ix, iy, w=iw)
    k = iw / 2620

    def at(px, py):
        return ix + (px - 580) * k, iy + (py - 168) * k

    for n, (px, py) in enumerate([(1290, 600), (2440, 380), (2470, 693), (1060, 1560)], 1):
        cx, cy = at(px, py)
        badge(s, cx - 0.15, cy - 0.15, str(n), d=0.3, size=11, outline=WHITE)
    notes = [
        ("1", "A SAM pops up on STK-01's route; old routes show dashed"),
        ("2", "Proposal: 2 strikes re-routed, 0 aircraft swapped (a fresh re-plan swaps 32)"),
        ("3", "The commander approves or rejects, 1.7 s after the alert"),
        ("4", "Synchronisation matrix: every sortie's timing, before and after"),
    ]
    ny = iy + ih + 0.12
    for i, (n, t) in enumerate(notes):
        y = ny + i * 0.26
        badge(s, ix, y + 0.02, n, d=0.2, size=8.5)
        text(s, ix + 0.3, y, iw - 0.3, 0.25, [t], size=11, color=INK)


STATIONS = [
    ("Something changes", "Fog forecast, an aircraft goes U/S, a pop-up SAM or new tasking",
     "event feed · fog model", "e.g. a SAM pops up on STK-01's route"),
    ("Check what is possible", "Which aircraft, crews, weapons and safe routes can still make each mission",
     "routing · fatigue model", "routes re-drawn around the new threat"),
    ("Find the smallest fix", "Most mission value, with a penalty on every change to the current plan",
     "Google OR-Tools CP-SAT", "2 strikes re-routed, 0 aircraft swapped"),
    ("Explain and double-check", "A reason for each change, an independent rule checker, 2,000 simulated days",
     "validator · Monte Carlo", "fulfilment stays at 95% (fresh re-plan: 32 swaps)"),
    ("Commander decides", "Approve: change orders go out. Reject: nothing moves",
     "map + timeline UI · audit log", "ready for a decision 1.7 s after the alert"),
]

TECH = [
    ("DECISION ENGINE", ["python", "google"], "Python 3.13 · Google OR-Tools CP-SAT · NumPy / SciPy"),
    ("FORECASTING", ["cloud"], "Open-Meteo weather model + airport METARs → fog model (logistic MOS)"),
    ("SERVICE", ["fastapi", "pydantic"], "FastAPI REST service · Pydantic data model"),
    ("INTERFACE", ["react", "typescript"], "React + TypeScript · deck.gl map, offline, India's official boundary"),
    ("COPILOT · HARDWARE", ["ollama", "laptop"], "Local open-weight LLM via Ollama (optional) · one laptop, no internet"),
]


def slide_technical(s):
    frame(s)
    pointer(s, 0.45, 1.2, 12.4, "Technologies to be used (e.g. programming languages, frameworks, hardware)")
    n, gap, x0 = len(TECH), 0.2, 0.45
    tw = (12.43 - gap * (n - 1)) / n
    for i, (label, icons, desc) in enumerate(TECH):
        x = x0 + i * (tw + gap)
        box(s, x, 1.58, tw, 1.02)
        text(s, x + 0.12, 1.66, tw - 0.8, 0.25, [{"runs": [(label, {"bold": True})]}], size=9.5, color=SAFF)
        for j, ic in enumerate(reversed(icons)):
            s.shapes.add_picture(str(ASSETS / "icons" / f"{ic}.png"), Inches(x + tw - 0.38 - j * 0.32),
                                 Inches(1.64), Inches(0.26), Inches(0.26))
        text(s, x + 0.12, 1.97, tw - 0.24, 0.6, [desc], size=10.5)

    pointer(s, 0.45, 2.78, 12.4, "Methodology and process for implementation (Flow Charts/Images/ working prototype)")

    # Inputs that feed the loop.
    chips = ["Aircraft status", "Crew rosters", "Weapons & stocks", "Weather", "Threat intel", "Airspace", "New tasking"]
    text(s, 0.45, 3.2, 0.75, 0.26, [{"runs": [("INPUTS", {"bold": True})]}], size=9.5, color=MUTED,
         anchor=MSO_ANCHOR.MIDDLE)
    cx = 1.25
    for c in chips:
        w = 0.16 + 0.075 * len(c)
        b = box(s, cx, 3.2, w, 0.26, fill=WHITE, line=RULE, radius=0.13)
        tf = b.text_frame
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, c, 9.5, INK)
        cx += w + 0.1

    n, gap = len(STATIONS), 0.34
    sw = (12.43 - gap * (n - 1)) / n
    sy, sh = 3.66, 1.72
    arrow(s, 0.45 + sw * 1.5 + gap, 3.48, 0.45 + sw * 1.5 + gap, sy - 0.02, color=BASE_GREY, width=1.25)
    for i, (title, desc, tag, eg) in enumerate(STATIONS):
        x = 0.45 + i * (sw + gap)
        core = i == 2
        box(s, x, sy, sw, sh, fill=INK if core else WHITE, line=None if core else RULE, radius=0.08)
        badge(s, x + 0.12, sy + 0.13, str(i + 1), d=0.34, size=12)
        text(s, x + 0.54, sy + 0.1, sw - 0.62, 0.42, [{"runs": [(title, {"bold": True})]}], size=12.5,
             color=WHITE if core else INK, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.12, sy + 0.6, sw - 0.24, 0.72, [desc], size=10.5,
             color=RGBColor(0xE4, 0xE9, 0xF0) if core else INK)
        text(s, x + 0.12, sy + sh - 0.33, sw - 0.24, 0.24, [{"runs": [(tag, {"bold": True})]}], size=9,
             color=RGBColor(0xF3, 0xB0, 0x83) if core else SAFF)
        if i < n - 1:
            arrow(s, x + sw + 0.04, sy + sh / 2, x + sw + gap - 0.04, sy + sh / 2, color=MUTED, width=1.75)
        text(s, x + 0.05, sy + sh + 0.1, sw - 0.1, 0.42, [{"runs": [(eg, {"italic": True})]}], size=9.5,
             color=MUTED)

    # The loop back: the next event goes round again.
    ly = 6.2
    x_first, x_last = 0.45 + sw / 2, 0.45 + 4 * (sw + gap) + sw / 2
    arrow(s, x_last, sy + sh + 0.55, x_last, ly, color=BASE_GREY, width=1.25, head=False)
    arrow(s, x_last, ly, x_first, ly, color=BASE_GREY, width=1.25, head=False)
    arrow(s, x_first, ly, x_first, sy + sh + 0.55, color=BASE_GREY, width=1.25)
    lab = text(s, 3.2, ly - 0.15, 6.9, 0.3, [{"runs": [
        ("The next event goes round the same loop. ", {"bold": True}),
        ("Missions already airborne stay frozen.", {})]}], size=10.5, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    lab.fill.solid()
    lab.fill.fore_color.rgb = WHITE
    text(s, 0.45, 6.5, 12.43, 0.28, [{"runs": [
        ("Built and working: ", {"bold": True, "color": GREEN}),
        ("55 automated engine tests · a browser test that clicks through the whole flow · "
         "every plan passes the independent rule checker", {})]}], size=10.5, align=PP_ALIGN.CENTER)


def slide_feasibility(s):
    frame(s)
    pointer(s, 0.45, 1.2, 6.2, "Analysis of the feasibility of the idea")

    cd = CategoryChartData()
    cd.categories = ["Western front", "Flood relief", "Earthquake relief"]
    cd.add_series("Manual-style planner", (85.1, 85.1, 79.2))
    cd.add_series("VAYU-SARTHI", (97.5, 96.3, 87.6))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45), Inches(1.55), Inches(6.1),
                            Inches(2.95), cd)
    ch = gf.chart
    ch.font.name, ch.font.size = FONT, Pt(10.5)
    ch.font.color.rgb = INK
    ch.has_title = True
    tp = ch.chart_title.text_frame.paragraphs[0]
    _run(tp, "Mission fulfilment, same fleet and rules (20 scenarios each)", 11, INK, bold=True)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 70, -8
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0.0"%"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(10.5), True
    for ser, col in zip(plot.series, (BASE_GREY, INK)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = 0, 110
    va.has_major_gridlines = False
    va.visible = False
    ca = ch.category_axis
    ca.format.line.color.rgb = RULE
    ca.tick_labels.font.size = Pt(10.5)

    stats = [("2.8 s", "to re-plan when the busiest base fogs in"),
             ("−71%", "aircraft moved vs a fresh re-plan (14.5 vs 49.5)"),
             ("70%", "of fog hours forecast (raw weather model: 26%)")]
    sw = (6.1 - 0.3) / 3
    for i, (big, small) in enumerate(stats):
        x = 0.45 + i * (sw + 0.15)
        text(s, x, 4.62, sw, 0.45, [{"runs": [(big, {"bold": True})]}], size=24, color=GREEN)
        text(s, x, 5.07, sw, 0.45, [small], size=10, color=MUTED)
    box(s, 0.45, 5.62, 6.1, 1.15)
    text(s, 0.6, 5.7, 5.85, 1.0, [
        {"runs": [("Already built and tested: ", {"bold": True}), ("engine, map, timeline, copilot; 55 tests", {})],
         "bullet": True, "after": 2},
        {"runs": [("Every plan passes an independent rule checker ", {}),
                  ("(crew rest, stocks, weather, airspace)", {"color": MUTED})], "bullet": True, "after": 2},
        {"runs": [("Open-source stack on one laptop, offline: ", {}), ("no licence or cloud cost", {"bold": True})],
         "bullet": True},
    ], size=11)

    # Right: the two remaining pointers as the column heads of one table.
    rows = [
        ["Potential challenges and risks", "Strategies for overcoming these challenges"],
        [[("Real data is classified; ", {"bold": True}), ("legacy systems such as IACCS", {})],
         "One adapter per data source behind one world model; notional data until access is granted"],
        [[("Commanders may not trust ", {"bold": True}), ("machine advice", {})],
         "A reason for every choice, a human approves every change, an independent checker (DRDO ETAI principles)"],
        [[("Theatre scale: ", {"bold": True}), ("100+ missions a day", {})],
         "~40 missions plan in 4–10 s today. Next: large-neighbourhood search, split by sector and time"],
        [[("Forecasts drift ", {"bold": True}), ("with season and region", {})],
         "Probabilities, not yes/no; tested on a held-out winter; retrained per region and season"],
        [[("Network cut off ", {"bold": True}), ("(degraded links)", {})],
         "Runs fully offline today. Next: a node at each base that syncs when the link returns"],
    ]
    table(s, 6.85, 1.22, [2.4, 3.63], rows, size=11, row_h=0.98, header_h=0.55, header_pointers=True)


AUDIENCE = [
    ("map", "Air tasking cell, Air HQ and Commands", "Seconds, not hours",
     "to re-plan after fog, an aircraft going U/S or a pop-up threat"),
    ("plane", "Squadrons and base operations", "−71% churn",
     "14.5 aircraft moved instead of 49.5 when a base fogs in"),
    ("crew", "Aircrew", "Rested crews",
     "rest, fatigue and night currency are hard rules in every plan"),
    ("lifebuoy", "NDMA and state disaster response", "+12.9 pts relief",
     "of requested flood-relief tonnage lifted by the same fleet"),
]

BENEFITS = [
    ("Operational", "+12.4 pts mission fulfilment from the same aircraft and crews"),
    ("Economic", "Open-source, one laptop, no licences; idle aircraft become ground spares at no cost"),
    ("Social", "Flood and earthquake relief planned rescues-first: 94.5% and 88.1% of relief lifted"),
    ("Environmental", "Every sortie carries a cost in the plan, so none is flown without need"),
    ("Strategic", "Indigenous and air-gapped: no foreign cloud, model or map service"),
]


QUAKE_W, QUAKE_H = 1920, 860  # tools/deck_assets/quake.jpg


def slide_impact(s):
    frame(s)
    pointer(s, 0.45, 1.2, 12.4, "Potential impact on the target audience")
    n, gap = len(AUDIENCE), 0.25
    cw = (12.43 - gap * (n - 1)) / n
    for i, (icon, who, big, small) in enumerate(AUDIENCE):
        x = 0.45 + i * (cw + gap)
        box(s, x, 1.58, cw, 1.92)
        circ = box(s, x + 0.15, 1.72, 0.52, 0.52, fill=SAFF_TINT, shape=MSO_SHAPE.OVAL)
        circ.name = f"icon-bg-{icon}"
        s.shapes.add_picture(str(ASSETS / "icons" / f"{icon}.png"), Inches(x + 0.26), Inches(1.83),
                             Inches(0.3), Inches(0.3))
        text(s, x + 0.8, 1.7, cw - 0.92, 0.56, [{"runs": [(who, {"bold": True})]}], size=11.5,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.15, 2.4, cw - 0.3, 0.4, [{"runs": [(big, {"bold": True})]}], size=19, color=GREEN)
        text(s, x + 0.15, 2.83, cw - 0.3, 0.6, [small], size=10.5, color=MUTED)

    pointer(s, 0.45, 3.72, 12.4, "Benefits of the solution (social, economic, environmental, etc.)")
    for i, (tag, line) in enumerate(BENEFITS):
        y = 4.12 + i * 0.52
        t = box(s, 0.45, y + 0.04, 1.4, 0.32, fill=SAFF_TINT, radius=0.16)
        tf = t.text_frame
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, tag, 10.5, SAFF, bold=True)
        text(s, 2.0, y, 4.75, 0.42, [line], size=11, anchor=MSO_ANCHOR.MIDDLE)

    iw = 5.6
    ih = iw * QUAKE_H / QUAKE_W
    picture(s, ASSETS / "quake.jpg", 12.88 - iw, 4.05, w=iw)
    text(s, 12.88 - iw, 4.05 + ih + 0.05, iw, 0.26, [{"runs": [
        ("Same engine, earthquake relief: ", {"bold": True}),
        ("thin air at 4,750 m rules out the medium helicopter", {})]}], size=10, color=MUTED)


REFS = [
    ("IAF AI decision-support tools for IACCS (IDRW)", "idrw.org/?p=383033", "https://idrw.org/?p=383033",
     "Where a tool like this would plug in"),
    ("DRDO ETAI framework for trustworthy AI in defence (2024)", "indiaai.gov.in",
     "https://indiaai.gov.in/news/trustworthy-ai-framework-launched-for-critical-defence-operations",
     "Human approval, reasons, independent checks"),
    ("Kessel Run KRADOS and Slapshot: ATO in the cloud (USAF)", "airandspaceforces.com",
     "https://www.airandspaceforces.com/afcent-can-now-generate-air-tasking-orders-in-the-cloud/",
     "What exists abroad; the gap is minimal-change retasking"),
    ("DARPA Adapting Cross-Domain Kill-Webs (ACK)", "darpa.mil",
     "https://www.darpa.mil/research/programs/adapting-cross-domain-kill-webs",
     "A decision aid to task and retask quickly"),
    ("Google OR-Tools, CP-SAT constraint solver", "developers.google.com/optimization",
     "https://developers.google.com/optimization", "The planning and retasking engine"),
    ("Glahn & Lowry (1972), Model Output Statistics, J. Appl. Meteor. 11", "doi.org/10.1175/1520-0450(1972)011",
     "https://doi.org/10.1175/1520-0450(1972)011%3C1203:TUOMOS%3E2.0.CO;2", "How the fog forecast is built"),
    ("Borbély (1982), two-process model of sleep; SAFTE in ICAO FRMS", "icao.int (FRMS modelling)",
     "https://www.icao.int/sites/default/files/sp-files/SAM/Documents/2012/FRMS11/FRMS%20GIG%20Modeling%20Presentation.pdf",
     "Crew fatigue and rest windows"),
    ("Minimal-deviation rescheduling under disruption (METU)", "avesis.metu.edu.tr",
     "https://avesis.metu.edu.tr/yayin/0895632c-7f5a-4e9b-abe3-de5753c8eec2/bi-objective-missile-rescheduling-for-a-naval-task-group-with-dynamic-disruptions",
     "A penalty on changing the current plan"),
    ("Open-Meteo forecasts · IEM METAR archive · Natural Earth (India view)",
     "open-meteo.com · mesonet.agron.iastate.edu", "https://open-meteo.com",
     "Real weather, airport observations, maps"),
    ("Flexible Use of Airspace in India (ICAO APAC IP08, 2023)", "icao.int",
     "https://www.icao.int/sites/default/files/APAC/Meetings/2023/2023%20SAIOSEACG%202/4-Information%20Papers/IP08-Benefits-of-Flexible-Use-of-Airspace-in-India.pdf",
     "Airspace that opens and closes in the day"),
]


def slide_references(s):
    frame(s)
    pointer(s, 0.45, 1.2, 12.4, "Details / Links of the reference and research work")
    rows = [["Reference", "Link", "What we took from it"]]
    for i, (ref, short, url, use) in enumerate(REFS, 1):
        rows.append([[(f"{i:>2}  ", {"color": SAFF, "bold": True}), (ref, {})],
                     [(short, {"link": url, "color": TX2})], use])
    table(s, 0.45, 1.6, [5.35, 3.4, 3.68], rows, size=10.5, row_h=0.42, header_h=0.36)
    b = box(s, 0.45, 6.3, 12.43, 0.48)
    live = [("Live prototype: ", {"bold": True}),
            (LIVE_URL.removeprefix("https://"), {"link": LIVE_URL, "color": TX2}), ("   ·   ", {"color": MUTED}),
            ("Code: ", {"bold": True})] if LIVE_URL else [("Our prototype (open source): ", {"bold": True})]
    text(s, 0.6, 6.32, 12.1, 0.44, [{"runs": live + [
        (REPO_URL.removeprefix("https://"), {"link": REPO_URL, "color": TX2}),
        ("   ·   all scenario data is notional; maps use India's official boundary", {"color": MUTED})]}],
         size=11 if not LIVE_URL else 10.5, anchor=MSO_ANCHOR.MIDDLE)
    b.name = "prototype-link"


def _navy_links(prs) -> None:
    """Hyperlinks in the template's navy instead of the theme's default blue."""
    theme = next(r.target_part for r in prs.slide_masters[0].part.rels.values() if r.reltype.endswith("/theme"))
    root = etree.fromstring(theme.blob)
    for tag in ("a:hlink", "a:folHlink"):
        el = root.find(f".//{qn(tag)}")
        if el is not None:
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, qn("a:srgbClr")).set("val", str(TX2))
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def main() -> None:
    global LIVE_URL, OUT
    ap = argparse.ArgumentParser()
    ap.add_argument("--live", default="", help="public URL of the deployed prototype (shown on slides 1 and 6)")
    ap.add_argument("out", nargs="?", type=Path, default=OUT)
    args = ap.parse_args()
    LIVE_URL, OUT = args.live.rstrip("/"), args.out
    prs = Presentation(str(ASSETS / "SIH2026-IDEA-Presentation-Format.pptx"))
    _navy_links(prs)
    slides = list(prs.slides)
    slide_title(slides[0])
    slide_solution(slides[1])
    slide_technical(slides[2])
    slide_feasibility(slides[3])
    slide_impact(slides[4])
    slide_references(slides[5])
    # The instructions slide may be deleted before upload (the template says so).
    sld_ids = prs.slides._sldIdLst
    last = list(sld_ids)[6]
    prs.part.drop_rel(last.get(qn("r:id")))
    sld_ids.remove(last)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
