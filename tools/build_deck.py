"""Build the SIH idea deck (PPTX) from real screenshots and measured results.

    python -I tools/build_deck.py [shots_dir] [out.pptx]

Screenshots come from the browser end-to-end run (frontend: `npm run e2e`, saved in frontend/e2e/shots).
The six slides follow the SIH idea template's sections (title, proposed solution, technical approach,
feasibility and viability, impact and benefits, research and references), so the content can be pasted
into the official template slide by slide. Numbers are the measured results in docs/PLAN.md section 8.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "frontend" / "e2e" / "shots"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "VAYU-SARTHI_SIH2026_idea.pptx"

FONT = "Arial"
BG = RGBColor(0x12, 0x12, 0x11)
PANEL = RGBColor(0x1E, 0x1E, 0x1C)
PANEL2 = RGBColor(0x2A, 0x2A, 0x28)
INK = RGBColor(0xFF, 0xFF, 0xFF)
INK2 = RGBColor(0xC3, 0xC2, 0xB7)
MUTED = RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x39, 0x87, 0xE5)
ACCENT_INK = RGBColor(0x86, 0xB6, 0xEF)
GOOD = RGBColor(0x0C, 0xA3, 0x0C)
BAND = RGBColor(0x1E, 0x6B, 0x1E)

# ---------- measured results (docs/PLAN.md section 8) ----------
R = {
    "fulfil": ("85.1%", "97.5%"),
    "ev": ("65.3%", "78.8%"),
    "churn": ("49.5", "14.5"),
    "retask_s": "2.8 s",
    "fog_pod": ("26%", "70%"),
    "bss": "+44%",
    "auc": "0.91",
    "p05": ("66.9%", "69.9%"),
    "hadr_fulfil": ("85.1%", "96.3%"),
    "hadr_t": ("81.6%", "94.5%"),
    "coa_loss": ("5.81", "3.01"),
    "tests": "30",
}


def rgb_text(run, size, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1, space_after=4):
    """paras: str | list of (str | list of (text, {size, color, bold, italic}))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(space_after)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, st in runs:
            r = para.add_run()
            r.text = t
            rgb_text(r, st.get("size", size), st.get("color", color), st.get("bold", bold), st.get("italic", False))
    return tb


def bullets(slide, x, y, w, h, items, size=13, color=INK2, gap=5, mark="▪", mark_color=ACCENT):
    paras = []
    for it in items:
        runs = [(f"{mark}  ", {"color": mark_color, "size": size})]
        runs += [(it, {})] if isinstance(it, str) else it
        paras.append(runs)
    return text(slide, x, y, w, h, paras, size=size, color=color, space_after=gap)


def box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)
    return c


def picture(slide, path: Path, x, y, w=None, h=None, crop=None, border=True):
    """Place a screenshot; crop=(left, top, right, bottom) fractions."""
    if not path.exists():
        box(slide, x, y, w or 4, h or 2.4, fill=PANEL2)
        text(slide, x, y, w or 4, h or 2.4, f"[{path.name}: run npm run e2e]", size=11, color=MUTED,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return None
    iw, ih = Image.open(path).size
    cl, ct, cr, cb = crop or (0, 0, 0, 0)
    aspect = (iw * (1 - cl - cr)) / (ih * (1 - ct - cb))
    if w is not None and h is None:
        h = w / aspect
    elif h is not None and w is None:
        w = h * aspect
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cl, ct, cr, cb
    if border:
        pic.line.color.rgb = RGBColor(0x4A, 0x49, 0x44)
        pic.line.width = Pt(0.75)
    return pic


def caption(slide, x, y, w, s):
    text(slide, x, y, w, 0.3, s, size=10, color=MUTED)


def frame(prs, title, kicker):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    band = box(s, 0, 0, 13.333, 0.22, fill=BAND, shape=MSO_SHAPE.RECTANGLE)
    band.text_frame.text = ""
    text(s, 0, 0.0, 13.333, 0.22, "UNCLASSIFIED // EXERCISE · NOTIONAL DATA", size=9, color=INK, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.5, 0.36, 9, 0.3, kicker.upper(), size=11, color=ACCENT_INK, bold=True)
    text(s, 0.5, 0.62, 12.3, 0.6, title, size=26, bold=True)
    text(s, 0.5, 7.1, 8, 0.3, "VAYU-SARTHI · SIH 2026 · PS 26250", size=9, color=MUTED)
    return s


def section(slide, x, y, w, label):
    text(slide, x, y, w, 0.3, label.upper(), size=11, color=MUTED, bold=True)


# ---------- slides ----------

def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    box(s, 0, 0, 13.333, 0.22, fill=BAND, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0, 0.0, 13.333, 0.22, "UNCLASSIFIED // EXERCISE · NOTIONAL DATA", size=9, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.6, 0.6, 6.4, 0.3, "SMART INDIA HACKATHON 2026 · IDEA PRESENTATION", size=12, color=ACCENT_INK, bold=True)
    text(s, 0.6, 1.0, 6.4, 0.9, "VAYU-SARTHI", size=48, bold=True)
    text(s, 0.6, 1.95, 6.2, 1.0, "An explainable decision engine that plans air operations, predicts disruption and "
         "re-plans in seconds, changing as little as possible. The commander decides.", size=17, color=INK2,
         spacing=1.15)
    rows = [
        ("Problem Statement ID", "26250", 0.4),
        ("Problem Statement title", "Air Power: Dynamic Air Operations & Resource Optimisation", 0.62),
        ("Organisation", "Ministry of Defence · Defence Services Staff College", 0.62),
        ("Theme", "Transportation & Logistics", 0.4),
        ("PS category", "Software", 0.4),
        ("Team ID", "________", 0.4),
        ("Team name", "________", 0.4),
        ("Demo video", "________", 0.4),
    ]
    y = 3.05
    for k, v, h in rows:
        text(s, 0.6, y, 2.3, h, k, size=12, color=MUTED)
        text(s, 2.9, y, 4.2, h, v, size=13, bold=True)
        y += h
    pic = picture(s, SHOTS / "01-plan.png", 7.35, 0.75, w=5.55)
    if pic is not None:
        caption(s, 7.35, 0.75 + pic.height / 914400 + 0.05, 5.6,
                "Working prototype: common operational picture, synchronisation matrix and KPIs for a 31-mission day.")
    box(s, 7.35, 4.6, 5.55, 2.2, fill=PANEL)
    text(s, 7.55, 4.7, 5.2, 0.3, "MEASURED ON THE PROTOTYPE", size=11, color=MUTED, bold=True)
    stats = [(f"+{float(R['fulfil'][1][:-1]) - float(R['fulfil'][0][:-1]):.1f} pts", "mission fulfilment vs a manual-style plan"),
             ("−71%", "aircraft reassigned on a retask vs re-planning from scratch"),
             (R["fog_pod"][1], f"of fog hours forecast (raw weather model: {R['fog_pod'][0]})")]
    yy = 5.05
    for big, small in stats:
        text(s, 7.55, yy, 1.6, 0.5, big, size=22, bold=True, color=INK)
        text(s, 9.2, yy + 0.08, 3.6, 0.5, small, size=12, color=INK2)
        yy += 0.55


def slide_solution(prs):
    s = frame(prs, "Plan, predict, re-plan in seconds, and explain every change", "Proposed solution")
    section(s, 0.5, 1.35, 6.8, "What it is")
    text(s, 0.5, 1.62, 6.9, 1.2, [[
        ("A decision-support system for the air staff. It fuses aircraft, crews, weapons, tankers, weather, threats "
         "and airspace into one picture, and computes the best air tasking plan. When something changes it proposes the ", {}),
        ("smallest set of changes", {"bold": True, "color": INK}),
        (" that recovers the most mission value, as a reviewable diff. A human approves every change.", {}),
    ]], size=12, color=INK2, spacing=1.1)
    section(s, 0.5, 2.85, 6.8, "How it addresses the problem statement")
    bullets(s, 0.5, 3.13, 6.9, 2.6, [
        [("Dynamic retasking: ", {"bold": True, "color": INK}),
         (f"churn-penalised re-optimisation; launched missions frozen; {R['churn'][1]} vs {R['churn'][0]} aircraft "
          f"reassigned for a fogged base, in {R['retask_s']}.", {})],
        [("Resource optimisation: ", {"bold": True, "color": INK}),
         ("one model across aircraft, crew fatigue, weapons, tankers, weather, threats and alert reserve.", {})],
        [("Weather: ", {"bold": True, "color": INK}),
         (f"fog forecast trained on real airport observations (Brier skill {R['bss']}), turned into probabilistic base closures.", {})],
        [("Uncertainty: ", {"bold": True, "color": INK}),
         ("2,000 simulated days, single points of failure, ground spares from idle aircraft.", {})],
        [("Commander's intent: ", {"bold": True, "color": INK}),
         ("three courses of action side by side, each trade in one sentence.", {})],
    ], size=11.5, gap=5)
    section(s, 0.5, 5.45, 6.8, "Innovation and uniqueness")
    bullets(s, 0.5, 5.73, 6.9, 1.15, [
        "A minimal-disruption diff, not a fresh plan · one integrated optimiser, not seven silos",
        "Why not, and what it would take, for every unplanned mission · intel-age-aware routing · COAs",
        "Air-gapped and indigenous · the same engine plans flood relief (HADR)",
    ], size=11.5, gap=3)
    pic = picture(s, SHOTS / "04-fog-proposal.png", 7.75, 1.4, w=5.05)
    if pic is not None:
        caption(s, 7.75, 1.4 + pic.height / 914400 + 0.04, 5.05,
                "Fog forecast → proposed retask: diff, naive re-plan comparison, KPI deltas, approve / reject.")
    pic2 = picture(s, SHOTS / "02b-coa-compare.png", 7.75, 4.85, w=5.05, crop=(0.15, 0.07, 0.15, 0.45))
    if pic2 is not None:
        caption(s, 7.75, 4.85 + pic2.height / 914400 + 0.04, 5.05,
                "Courses of action: one situation, three intents, each trade in one line.")


def slide_technical(prs):
    s = frame(prs, "Technical approach: an optimiser with a dashboard on top", "Technical approach")
    steps = [
        ("Observe", "Data adapters", "aircraft · crews · stocks · met · intel · airspace"),
        ("Fuse", "World state", "one typed model, timestamps, notional data"),
        ("Predict", "Predictive layer", "fog MOS · P(serviceable) · crew fatigue · intel age"),
        ("Decide", "Optimiser", "threat-aware routing · CP-SAT allocation · minimal-disruption retask · COAs · spares"),
        ("Explain", "Explain + verify", "reason codes · independent validator · Monte Carlo"),
        ("Act", "Commander", "diff + KPIs → approve / reject → ATO change orders, audit log"),
    ]
    x, w, gap = 0.5, 1.95, 0.17
    for i, (k, t, d) in enumerate(steps):
        bx = x + i * (w + gap)
        box(s, bx, 1.45, w, 1.6, fill=PANEL2 if k != "Decide" else RGBColor(0x17, 0x33, 0x55))
        text(s, bx + 0.12, 1.52, w - 0.2, 0.3, k.upper(), size=10, color=ACCENT_INK, bold=True)
        text(s, bx + 0.12, 1.78, w - 0.2, 0.35, t, size=14, bold=True)
        text(s, bx + 0.12, 2.15, w - 0.22, 0.9, d, size=10.5, color=INK2, spacing=1.05)
        if i < len(steps) - 1:
            arrow(s, bx + w + 0.01, 2.25, bx + w + gap - 0.01, 2.25, color=INK2)
    text(s, 0.5, 3.1, 12.3, 0.3, "↺  Events (weather, unserviceability, pop-up threats, new tasking) re-enter the loop; "
         "every proposal waits for a human decision.", size=11, color=MUTED)

    section(s, 0.5, 3.55, 6.0, "Technologies")
    rows = [
        ("Decision engine", "Python 3.13 · Google OR-Tools CP-SAT · NumPy / SciPy (Dijkstra) · Pydantic"),
        ("Service", "FastAPI (REST); one process serves API and UI; runs offline on one laptop"),
        ("User interface", "React 19 + TypeScript · deck.gl map with an offline basemap (India's official boundary) · Zustand"),
        ("Prediction", "Logistic MOS on Open-Meteo NWP, labelled with IEM METAR; held-out winter validation"),
        ("Assurance", f"{R['tests']} engine tests · independent constraint validator · Playwright browser end-to-end test"),
    ]
    y = 3.85
    for k, v in rows:
        text(s, 0.5, y, 1.6, 0.5, k, size=11.5, color=INK, bold=True)
        text(s, 2.1, y, 4.6, 0.6, v, size=11.5, color=INK2, spacing=1.05)
        y += 0.6

    section(s, 7.0, 3.55, 5.8, "Key methods")
    bullets(s, 7.0, 3.85, 5.85, 3.2, [
        "CP-SAT with optional intervals for aircraft, crews and tankers; cumulative alert-reserve floor; "
        "time-on-target domains from crew rest, night currency and base closures",
        "Retask objective rewards keeping what is planned, weighted by time to launch (×1 planned, ×3 briefed, "
        "×8 armed); launched missions are frozen",
        "Threat hazard field with intel-age envelope growth and SEAD suppression; Dijkstra plus any-angle smoothing",
        "Two-process fatigue model gives each crew's fit time windows",
        "One mission-success model (serviceability with spares, tankers, SEAD before strike, ingress risk) "
        "drives both the expected-value KPI and the Monte Carlo",
    ], size=11.5, gap=5)


def slide_feasibility(prs):
    s = frame(prs, "Feasibility and viability: the prototype already works", "Feasibility and viability")
    section(s, 0.5, 1.35, 6.4, "Measured results (notional scenarios, laptop CPU)")
    rows = [
        ("Metric", "Baseline", "VAYU-SARTHI"),
        ("Priority-weighted fulfilment · 20 scenarios", f"{R['fulfil'][0]} manual-style", R["fulfil"][1]),
        ("Expected mission value on the day", R["ev"][0], R["ev"][1]),
        ("Aircraft reassigned, base fogged · 8 scenarios", f"{R['churn'][0]} naive re-plan", f"{R['churn'][1]} in {R['retask_s']}"),
        ("Fog hours forecast · held-out winter", f"{R['fog_pod'][0]} raw model", f"{R['fog_pod'][1]} (AUC {R['auc']})"),
        ("Bad-day (p05) success · ground spares", R["p05"][0], f"{R['p05'][1]}, 0 flying changes"),
        ("Expected losses · Min-risk COA", f"{R['coa_loss'][0]} Max effect", R["coa_loss"][1]),
        ("HADR relief tonnage · 20 scenarios", f"{R['hadr_t'][0]} manual-style", R["hadr_t"][1]),
    ]
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.7), Inches(6.5), Inches(0.38 * len(rows))).table
    tbl._tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId").text = \
        "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tbl.columns[0].width = Inches(3.15)
    tbl.columns[1].width = Inches(1.65)
    tbl.columns[2].width = Inches(1.7)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = PANEL2 if i == 0 else (PANEL if i % 2 else BG)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            p = c.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = v
            rgb_text(r, 10.5 if i else 10, INK if (j == 2 or i == 0) else INK2, bold=(i == 0 or j == 2))
    text(s, 0.5, 1.75 + 0.38 * len(rows), 6.5, 0.6,
         "Every plan in the tests and benchmarks passes an independent constraint checker (crew rest, stocks, "
         "reserve, weather, airspace). Plans in 4–10 s for 35–40 missions; retasks in 1–3 s.", size=10.5, color=MUTED)

    section(s, 7.35, 1.35, 5.6, "Risks and how we handle them")
    risks = [
        ("Classified data, legacy systems", "Adapter per source behind one world model; notional data until access; IACCS/AFNet are adapters, not a rewrite."),
        ("Trust in machine advice", "Nothing changes without approval; a reason for every decision; independent validator; DRDO ETAI principles."),
        ("Theatre scale", "Large-neighbourhood search and decomposition by sector and time; retask touches only the affected missions."),
        ("Model drift (weather, maintenance)", "Probabilities, not yes/no; calibrated and verified; retrained per region and season."),
        ("Degraded links (DDIL)", "Runs air-gapped; an edge node per base syncs on reconnect."),
    ]
    y = 1.7
    for k, v in risks:
        box(s, 7.35, y, 5.5, 0.92, fill=PANEL)
        text(s, 7.5, y + 0.07, 5.25, 0.3, k, size=12, bold=True)
        text(s, 7.5, y + 0.36, 5.25, 0.6, v, size=10.5, color=INK2, spacing=1.05)
        y += 1.02


def slide_impact(prs):
    s = frame(prs, "Impact and benefits", "Impact and benefits")
    cols = [
        ("Operational", [
            "Re-planning goes from hours to seconds, and the plan stays stable: one third of the changes",
            f"More mission value from the same fleet: +{float(R['fulfil'][1][:-1]) - float(R['fulfil'][0][:-1]):.1f} pts fulfilment",
            "Safer: crew fatigue, night currency and threat risk are constraints, not afterthoughts",
            "A known bad day: what fails, why, and which aircraft the plan leans on",
        ]),
        ("Resources and cost", [
            "Sorties, tanker sorties and guided weapons are used only where they buy effect",
            "Commander's intent trades effect for losses and stocks explicitly (Min risk: half the expected losses)",
            "Idle aircraft become ground spares at no cost to the plan",
        ]),
        ("National and dual-use", [
            f"Same engine for flood relief: {R['hadr_t'][0]} → {R['hadr_t'][1]} of relief tonnage with the same fleet",
            "Rescues first; helicopters where there is no runway; flights stay in Indian airspace",
            "Indigenous and air-gapped: no foreign cloud or map service; runs on one laptop",
        ]),
    ]
    for i, (h, items) in enumerate(cols):
        x = 0.5 + i * 4.18
        box(s, x, 1.4, 3.98, 2.95, fill=PANEL)
        text(s, x + 0.18, 1.5, 3.6, 0.35, h, size=15, bold=True)
        bullets(s, x + 0.18, 1.9, 3.65, 2.4, items, size=11.5, gap=5)
    text(s, 0.5, 4.45, 12.3, 0.35, [[("Who uses it: ", {"bold": True, "color": INK}),
                                      ("air tasking order cells at Air HQ and Commands, base operations, air defence controllers, "
                                       "and NDMA / state disaster-response coordination during HADR.", {})]],
         size=12, color=INK2)
    shots = [("02d-robustness.png", (0.15, 0.07, 0.15, 0.03), "Robustness: 2,000 simulated days, with and without spares"),
             ("11-hadr-breach.png", (0, 0, 0, 0), "Flood relief: a breach becomes a P10 rescue"),
             ("05-fog-proposal-aircraft-view.png", (0, 0.55, 0.42, 0), "Aircraft lanes: closures, turnarounds, changes")]
    for i, (f, crop, cap) in enumerate(shots):
        x = 0.5 + i * 4.18
        pic = picture(s, SHOTS / f, x, 4.9, w=3.98, crop=crop)
        if pic is not None and pic.height / 914400 > 1.85:
            pic.height = Inches(1.85)
            iw, ih = Image.open(SHOTS / f).size
            # keep the aspect by trimming the bottom instead of squashing
            visible_h = 1.85 / 3.98 * iw * (1 - crop[0] - crop[2])
            pic.crop_bottom = max(0.0, 1 - crop[1] - visible_h / ih)
        caption(s, x, 6.8, 3.98, cap)


def slide_references(prs):
    s = frame(prs, "Research and references", "Research and references")
    left = [
        "Problem and context",
        "IAF AI decision-support tools for IACCS (IDRW, idrw.org/?p=383033)",
        "DRDO ETAI framework for trustworthy AI in defence (IndiaAI, 2024)",
        "DARPA Adapting Cross-Domain Kill-Webs (darpa.mil)",
        "Kessel Run: KRADOS and Slapshot, ATO in the cloud (Air & Space Forces Magazine)",
        "Flexible Use of Airspace in India (ICAO APAC IP08, 2023)",
        "",
        "Methods",
        "Perron & Didier, CP-SAT / Google OR-Tools (developers.google.com/optimization)",
        "Glahn & Lowry (1972), Model Output Statistics, J. Applied Meteorology 11",
        "Borbély (1982), two-process model of sleep regulation, Human Neurobiology 1",
        "SAFTE fatigue model overview (ICAO FRMS)",
        "Minimal-deviation rescheduling under disruption (METU, naval task groups)",
    ]
    right = [
        "Data (public)",
        "Open-Meteo historical forecast and forecast APIs (open-meteo.com)",
        "Iowa Environmental Mesonet METAR archive (mesonet.agron.iastate.edu)",
        "Natural Earth boundaries, India point of view (naturalearthdata.com)",
        "Airfield and district headquarters coordinates (public)",
        "",
        "This submission",
        "Working prototype: engine, map, synchronisation matrix, retask review, why-not and what-it-would-take, fog forecast, COAs, robustness, HADR",
        "Design and roadmap: docs/PLAN.md · all scenario data notional",
    ]

    def col(x, items):
        paras = []
        for it in items:
            if it in ("Problem and context", "Methods", "Data (public)", "This submission"):
                paras.append([(it.upper(), {"size": 11, "color": MUTED, "bold": True})])
            elif it == "":
                paras.append([(" ", {"size": 6})])
            else:
                paras.append([("▪  ", {"color": ACCENT}), (it, {})])
        text(s, x, 1.4, 6.0, 5.6, paras, size=12, color=INK2, space_after=5)

    col(0.5, left)
    col(6.9, right)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_title(prs)
    slide_solution(prs)
    slide_technical(prs)
    slide_feasibility(prs)
    slide_impact(prs)
    slide_references(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
